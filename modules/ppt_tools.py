import os
import fitz  # PyMuPDF
from pptx import Presentation
from pptx.util import Inches

def convert_pdf_to_ppt(files, out_dir, progress_callback=None):
    """
    Mengonversi halaman PDF menjadi slide PowerPoint (.pptx).
    Setiap halaman PDF di-render menjadi gambar HD dan dimasukkan ke dalam slide PPTX.
    """
    success_count = 0
    error_msg = ""
    
    for file_path in files:
        try:
            base_name = os.path.splitext(os.path.basename(file_path))[0]
            output_path = os.path.join(out_dir, f"{base_name}.pptx")

            doc = fitz.open(file_path)
            prs = Presentation()

            # Sesuaikan rasio slide dengan rasio halaman pertama PDF
            if len(doc) > 0:
                first_page = doc[0]
                rect = first_page.rect
                prs.slide_width = Inches(rect.width / 72)
                prs.slide_height = Inches(rect.height / 72)

            blank_layout = prs.slide_layouts[6]  # Layout slide kosong

            for page_num in range(len(doc)):
                page = doc[page_num]
                pix = page.get_pixmap(dpi=200)  # Render resolusi tinggi
                
                temp_img_path = os.path.join(out_dir, f"_temp_p{page_num}.png")
                pix.save(temp_img_path)

                slide = prs.slides.add_slide(blank_layout)
                slide.shapes.add_picture(temp_img_path, 0, 0, prs.slide_width, prs.slide_height)

                if os.path.exists(temp_img_path):
                    os.remove(temp_img_path)

            doc.close()
            prs.save(output_path)
            success_count += 1
        except Exception as e:
            error_msg = str(e)

    return success_count, len(files), error_msg


def convert_ppt_to_pdf(files, out_dir, progress_callback=None):
    """
    Mengonversi file PPTX/PPT menjadi PDF menggunakan MS Office PowerPoint Interop.
    """
    success_count = 0
    error_msg = ""
    
    try:
        import win32com.client
        has_win32 = True
    except ImportError:
        has_win32 = False

    for file_path in files:
        try:
            base_name = os.path.splitext(os.path.basename(file_path))[0]
            output_path = os.path.join(out_dir, f"{base_name}.pdf")
            abs_file = os.path.abspath(file_path)
            abs_out = os.path.abspath(output_path)

            if has_win32:
                ppt_app = win32com.client.Dispatch("PowerPoint.Application")
                presentation = ppt_app.Presentations.Open(abs_file, WithWindow=False)
                presentation.SaveAs(abs_out, 32)  # 32 = Format Export PDF
                presentation.Close()
                ppt_app.Quit()
                success_count += 1
            else:
                error_msg = "Library 'pywin32' belum terinstall atau MS PowerPoint tidak tersedia."
        except Exception as e:
            error_msg = str(e)

    return success_count, len(files), error_msg